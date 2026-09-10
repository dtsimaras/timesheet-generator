"""
v0.1 -- Δημιουργεί το μηνιαίο Timesheet (.xlsx) στο ακριβές layout που ζητάει
το IBM template: τίτλοι, sums, γκριζαρισμένα ΣΚ/αργίες, dropdown στα Comments.

Εγκατάσταση (μία φορά):
    pip3 install openpyxl holidays

Εκτέλεση:
    python3 main.py
"""

import calendar
from datetime import date

import holidays
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

# ---------------------------------------------------------------------------
# ΒΗΜΑ 1: ΡΥΘΜΙΣΕΙΣ -- ό,τι αλλάζει κάθε μήνα ή ανά χρήστη μπαίνει εδώ
# ---------------------------------------------------------------------------
YEAR = 2026
MONTH = 9

fullname = "Tsimaras Dimitrios"  # TODO: βάλε το δικό σου, ακριβώς όπως πρέπει να φαίνεται

DEFAULT_HOURS = 8
OUTPUT_FILENAME = f"Project_Timesheet_{fullname.replace(' ', '_')}.xlsx"

# Οι στήλες με τη σειρά που τις θέλει το template (8 στήλες, A ως H)
HEADERS = [
    "Date", "Name", "Days", "Hours of Services",
    "Overtimes", "StandBy by Day", "Comments", "Client Name",
]
NUM_COLUMNS = len(HEADERS)

# Οι επιλογές του dropdown στη στήλη "Comments"
COMMENT_OPTIONS = [
    "Annual Leave", "Sick Leave", "Student Leave",
    "Parental Leave", "StandBy Servises",
]

# "Στυλ" που ξαναχρησιμοποιούμε παρακάτω σε πολλά κελιά
HEADER_FILL = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
GREY_FILL = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")

# ---------------------------------------------------------------------------
# ΒΗΜΑ 2: ΑΡΓΙΕΣ ΕΛΛΑΔΑΣ -- ίδιο με πριν, υπολογίζονται αυτόματα
# ---------------------------------------------------------------------------
greek_holidays = holidays.Greece(years=YEAR)
_, days_in_month = calendar.monthrange(YEAR, MONTH)

# Οι γραμμές δεδομένων ξεκινούν στη γραμμή 6 (μετά τους 3 τίτλους + τα sums
# + τα headers). Το κρατάμε σε μεταβλητή γιατί το χρειαζόμαστε παρακάτω
# τρεις φορές (formulas, validation, loop) -- αν αλλάξει το layout, αλλάζει
# μόνο εδώ.
DATA_START_ROW = 6
DATA_END_ROW = DATA_START_ROW + days_in_month - 1

wb = Workbook()
ws = wb.active
ws.title = "Timesheet"

# Το τελευταίο γράμμα στήλης (H, αφού έχουμε 8 στήλες) -- το χρειαζόμαστε
# για τα merge_cells παρακάτω. chr(ord('A') + 7) == 'H'
last_col_letter = chr(ord("A") + NUM_COLUMNS - 1)

# ---------------------------------------------------------------------------
# ΒΗΜΑ 3: ΟΙ 3 ΤΙΤΛΟΙ (γραμμές 1-3), merged σε όλο το πλάτος
# ---------------------------------------------------------------------------
# merge_cells ενώνει ένα εύρος κελιών σε ένα -- ΠΡΟΣΟΧΗ: γράφεις τιμή ΜΟΝΟ
# στο πάνω-αριστερά κελί του εύρους (π.χ. A1), ποτέ στα υπόλοιπα.
ws.merge_cells(f"A1:{last_col_letter}1")
ws["A1"] = "IBM"

ws.merge_cells(f"A2:{last_col_letter}2")
ws["A2"] = "Monthly Consumption"

ws.merge_cells(f"A3:{last_col_letter}3")
ws["A3"] = f"1/{MONTH}/{YEAR} - {days_in_month}/{MONTH}/{YEAR}"

for row in (1, 2, 3):
    cell = ws[f"A{row}"]
    cell.font = Font(name="Arial", bold=True, size=12 if row == 1 else 11)
    cell.alignment = Alignment(horizontal="center")

# ---------------------------------------------------------------------------
# ΒΗΜΑ 4: ΓΡΑΜΜΗ 4 -- τα 3 sums πάνω από Days+Hours / Overtimes / StandBy
# ---------------------------------------------------------------------------
# ΣΗΜΑΝΤΙΚΟ: αυτά ΔΕΝ είναι νούμερα υπολογισμένα από την Python -- είναι
# πραγματικοί Excel τύποι (strings που ξεκινάνε με "="). Έτσι, αν αργότερα
# αλλάξεις χειροκίνητα ώρες μέσα στο Excel, το sum ενημερώνεται μόνο του.
ws.merge_cells("C4:D4")
ws["C4"] = f"=SUM(C{DATA_START_ROW}:C{DATA_END_ROW})"
ws["E4"] = f"=SUM(E{DATA_START_ROW}:E{DATA_END_ROW})"
ws["F4"] = f"=SUM(F{DATA_START_ROW}:F{DATA_END_ROW})"

for cell_ref in ("C4", "E4", "F4"):
    cell = ws[cell_ref]
    cell.number_format = "0.00"
    cell.alignment = Alignment(horizontal="center")

for col in range(1, NUM_COLUMNS + 1):
    ws.cell(row=4, column=col).fill = HEADER_FILL

# ---------------------------------------------------------------------------
# ΒΗΜΑ 5: ΓΡΑΜΜΗ 5 -- τα headers των στηλών
# ---------------------------------------------------------------------------
for col, header in enumerate(HEADERS, start=1):
    cell = ws.cell(row=5, column=col, value=header)
    cell.font = Font(name="Arial", bold=True)
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(horizontal="center")

# ---------------------------------------------------------------------------
# ΒΗΜΑ 6: DROPDOWN στη στήλη "Comments" (G) -- v0.1: μένει άδειο,
# απλά έτοιμο για χειροκίνητη επιλογή (το v0.2 θα το γεμίζει αυτόματα)
# ---------------------------------------------------------------------------
dv = DataValidation(
    type="list",
    formula1='"' + ",".join(COMMENT_OPTIONS) + '"',
    allow_blank=True,
)
ws.add_data_validation(dv)
dv.add(f"G{DATA_START_ROW}:G{DATA_END_ROW}")

# ---------------------------------------------------------------------------
# ΒΗΜΑ 7: ΓΡΑΜΜΕΣ ΔΕΔΟΜΕΝΩΝ -- μία ανά ημερολογιακή μέρα του μήνα
# ---------------------------------------------------------------------------
for day_num in range(1, days_in_month + 1):
    row = DATA_START_ROW + day_num - 1
    current_date = date(YEAR, MONTH, day_num)
    weekday_index = current_date.weekday()  # 0=Δευτέρα ... 6=Κυριακή
    is_weekend = weekday_index >= 5
    is_holiday = current_date in greek_holidays
    is_off_day = is_weekend or is_holiday

    date_cell = ws.cell(row=row, column=1, value=current_date)
    date_cell.number_format = "ddd d-mmm"  # εμφανίζεται σαν "Tue 1-Sep"

    if not is_off_day:
        ws.cell(row=row, column=2, value=fullname)
        ws.cell(row=row, column=3, value=1)
        hours_cell = ws.cell(row=row, column=4, value=DEFAULT_HOURS)
        hours_cell.number_format = "0.00"

    if is_off_day:
        for col in range(1, NUM_COLUMNS + 1):
            ws.cell(row=row, column=col).fill = GREY_FILL
        if is_holiday:
            # Σχόλιο (κίτρινο post-it πάνω στο κελί) ώστε να θυμάσαι ΓΙΑΤΙ
            # ήταν γκρι η μέρα -- δεν πειράζει τη στήλη Comments/dropdown.
            holiday_name = greek_holidays.get(current_date)
            date_cell.comment = Comment(holiday_name, "main.py")

# ---------------------------------------------------------------------------
# ΒΗΜΑ 8: ΠΛΑΤΗ ΣΤΗΛΩΝ & ΑΠΟΘΗΚΕΥΣΗ
# ---------------------------------------------------------------------------
widths = {"A": 12, "B": 24, "C": 8, "D": 16, "E": 12, "F": 14, "G": 18, "H": 16}
for col_letter, width in widths.items():
    ws.column_dimensions[col_letter].width = width

wb.save(OUTPUT_FILENAME)
print(f"Έτοιμο: {OUTPUT_FILENAME}")


# coloring seems off and merged cells and total days hour cells. We also need to fix the borders. The 3 rows of title shouldnt have any lines between them, Every other cell should show all over the table and the whole table surreounder by bolder black