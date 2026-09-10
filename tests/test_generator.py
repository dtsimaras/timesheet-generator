import unittest
from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from timesheet.config import TimesheetConfig
from timesheet.generator import (
    generate_timesheet,
    output_filename,
    reporting_dates,
    reporting_output_directory,
    selected_week_dates,
)


class ReportingDatesTests(unittest.TestCase):
    def test_week_one_selects_the_first_calendar_week_without_hiding_dates(
        self,
    ) -> None:
        month_dates = reporting_dates(self.config(week=1))
        dates = selected_week_dates(self.config(week=1), month_dates)

        self.assertEqual(len(month_dates), 30)
        self.assertEqual(dates, {
            date(2026, 9, 1), date(2026, 9, 2), date(2026, 9, 3),
            date(2026, 9, 4), date(2026, 9, 5), date(2026, 9, 6),
        })

    def test_full_month_contains_every_day(self) -> None:
        dates = reporting_dates(self.config(week=None))

        self.assertEqual(len(dates), 30)
        self.assertEqual(dates[0], date(2026, 9, 1))
        self.assertEqual(dates[-1], date(2026, 9, 30))

    def test_output_uses_the_reporting_folders_and_stable_filename(self) -> None:
        config = self.config(week=2)

        folder = reporting_output_directory(config, Path("output-root"))

        self.assertEqual(folder, Path("output-root/output/2026/09/week-2"))
        self.assertEqual(
            output_filename(config), "Project_Timesheet_Example_Developer.xlsx"
        )

    def test_week_selection_prefills_only_that_week_in_a_full_month_workbook(
        self,
    ) -> None:
        with TemporaryDirectory() as directory:
            output_path = generate_timesheet(self.config(week=1), Path(directory))
            sheet = load_workbook(output_path, data_only=False)["Timesheet"]

        self.assertEqual(sheet.max_row, 35)
        self.assertEqual(sheet["A2"].value, "Monthly Consumption")
        self.assertEqual(sheet["A3"].value, "1/9/2026 - 30/9/2026")
        self.assertEqual(sheet["B6"].value, "Example Developer")
        self.assertEqual(sheet["D6"].value, 8)
        self.assertIsNone(sheet["B12"].value)
        self.assertIsNone(sheet["D12"].value)
        self.assertEqual(sheet["C4"].value, "=SUM(C6:C35)")
        self.assertEqual(sheet["D4"].value, "=SUM(D6:D35)")
        self.assertEqual(sheet["H4"].fill.fgColor.rgb, "00DCE6F1")
        self.assertEqual(sheet["A4"].border.top.style, "thin")

    def test_greek_public_holiday_is_never_prefilled(self) -> None:
        with TemporaryDirectory() as directory:
            output_path = generate_timesheet(
                self.config(week=None, month=10), Path(directory)
            )
            sheet = load_workbook(output_path, data_only=False)["Timesheet"]

        self.assertEqual(sheet["A33"].value.date(), date(2026, 10, 28))
        self.assertIsNone(sheet["B33"].value)
        self.assertIsNone(sheet["D33"].value)
        self.assertEqual(sheet["A33"].fill.fgColor.rgb, "00D9D9D9")

    @staticmethod
    def config(week: int | None, month: int = 9) -> TimesheetConfig:
        return TimesheetConfig(
            employee_name="Example Developer",
            year=2026,
            month=month,
            week=week,
            default_hours=8,
            company_name="Example Company",
            holiday_country="GR",
            output_directory="output",
        )


if __name__ == "__main__":
    unittest.main()
