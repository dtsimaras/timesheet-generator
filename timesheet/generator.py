"""Build the Timesheet worksheet with openpyxl."""

import calendar
import re
from datetime import date, timedelta
from pathlib import Path

import holidays
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from timesheet.config import TimesheetConfig

HEADERS = [
    "Date", "Name", "Days", "Hours of Services", "Overtimes",
    "StandBy by Day", "Comments", "Client Name",
]
COMMENT_OPTIONS = [
    "Annual Leave", "Sick Leave", "Student Leave", "Parental Leave", "StandBy Services",
]

HEADER_FILL = PatternFill("solid", fgColor="DCE6F1")
OFF_DAY_FILL = PatternFill("solid", fgColor="D9D9D9")
THIN_SIDE = Side(style="thin", color="000000")
THIN_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)


def generate_timesheet(config: TimesheetConfig, project_root: Path) -> Path:
    """Create a workbook and return the generated file path."""
    dates = reporting_dates(config)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Timesheet"
    sheet.sheet_view.showGridLines = False

    write_titles(sheet, config, dates)
    write_summary_and_headers(sheet, len(dates))
    write_daily_rows(sheet, config, dates)
    apply_table_borders(sheet, last_row=5 + len(dates))
    set_column_widths(sheet)

    output_dir = reporting_output_directory(config, project_root)
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / output_filename(config)
    workbook.save(output_path)
    return output_path


def reporting_dates(config: TimesheetConfig) -> list[date]:
    """Return every date in the selected month.

    The week setting controls which workdays receive default values; it does
    not remove the other dates from the monthly workbook.
    """
    _, days_in_month = calendar.monthrange(config.year, config.month)
    return [date(config.year, config.month, day) for day in range(1, days_in_month + 1)]

def selected_week_dates(config: TimesheetConfig, month_dates: list[date]) -> set[date]:
    """Return dates that should receive default name, day, and hour values."""
    if config.week is None:
        return set(month_dates)

    first_day = month_dates[0]
    first_week_monday = first_day - timedelta(days=first_day.weekday())
    week_start = first_week_monday + timedelta(weeks=config.week - 1)
    week_end = week_start + timedelta(days=6)
    selected_dates = {day for day in month_dates if week_start <= day <= week_end}
    if not selected_dates:
        raise ValueError("config.json: the selected week does not overlap this month.")
    return selected_dates


def write_titles(sheet, config: TimesheetConfig, dates: list[date]) -> None:
    """Write the three clean title rows at the top of the worksheet."""
    sheet.merge_cells("A1:H1")
    sheet.merge_cells("A2:H2")
    sheet.merge_cells("A3:H3")
    sheet["A1"] = config.company_name
    sheet["A2"] = "Monthly Consumption"
    start_date = format_date(dates[0])
    end_date = format_date(dates[-1])
    sheet["A3"] = f"{start_date} - {end_date}"

    for row, size in ((1, 12), (2, 11), (3, 11)):
        cell = sheet[f"A{row}"]
        cell.font = Font(name="Arial", bold=True, size=size)
        cell.alignment = Alignment(horizontal="center")


def write_summary_and_headers(sheet, days_count: int) -> None:
    """Write totals, blue summary cells, and the column headers."""
    data_start_row = 6
    data_end_row = data_start_row + days_count - 1

    for column in "ABCDEFGH":
        sheet[f"{column}4"].fill = HEADER_FILL

    for column in ("C", "D", "E", "F"):
        cell = sheet[f"{column}4"]
        cell.value = f"=SUM({column}{data_start_row}:{column}{data_end_row})"
        cell.number_format = "0.00"
        cell.alignment = Alignment(horizontal="center")

    for column, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=5, column=column, value=header)
        cell.font = Font(name="Arial", bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")


def write_daily_rows(sheet, config: TimesheetConfig, dates: list[date]) -> None:
    """Fill normal workdays, shade non-working days, and add the comments list."""
    public_holidays = holidays.country_holidays(
        config.holiday_country, years=config.year
    )
    data_start_row = 6
    dates_to_fill = selected_week_dates(config, dates)

    for row, current_date in enumerate(dates, start=data_start_row):
        date_cell = sheet.cell(row=row, column=1, value=current_date)
        date_cell.number_format = "ddd d-mmm"
        is_weekend = current_date.weekday() >= 5
        is_holiday = current_date in public_holidays

        if is_weekend or is_holiday:
            for column in range(1, len(HEADERS) + 1):
                sheet.cell(row=row, column=column).fill = OFF_DAY_FILL
            if is_holiday:
                date_cell.comment = Comment(
                    public_holidays.get(current_date), "Timesheet generator"
                )
            continue

        if current_date not in dates_to_fill:
            continue

        sheet.cell(row=row, column=2, value=config.employee_name)
        sheet.cell(row=row, column=3, value=1)
        hours_cell = sheet.cell(row=row, column=4, value=config.default_hours)
        hours_cell.number_format = "0.00"

    validation = DataValidation(
        type="list",
        formula1='"' + ",".join(COMMENT_OPTIONS) + '"',
        allow_blank=True,
    )
    sheet.add_data_validation(validation)
    validation.add(f"G{data_start_row}:G{data_start_row + len(dates) - 1}")


def apply_table_borders(sheet, last_row: int) -> None:
    """Use a thin border throughout the table and no heavy outer frame."""
    for row in sheet.iter_rows(
        min_row=4, max_row=last_row, min_col=1, max_col=len(HEADERS)
    ):
        for cell in row:
            cell.border = THIN_BORDER


def set_column_widths(sheet) -> None:
    """Give the template's eight columns readable widths."""
    widths = {"A": 12, "B": 24, "C": 8, "D": 18, "E": 12, "F": 16, "G": 18, "H": 16}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


def output_filename(config: TimesheetConfig) -> str:
    """Keep the workbook filename stable; the folder identifies its period."""
    safe_name = "_".join(config.employee_name.split())
    safe_name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", safe_name).strip(". ")
    if not safe_name:
        safe_name = "Timesheet"
    return f"Project_Timesheet_{safe_name}.xlsx"


def reporting_output_directory(config: TimesheetConfig, project_root: Path) -> Path:
    """Store each generated workbook under output/year/month/week."""
    week_folder = f"week-{config.week}" if config.week is not None else "all-weeks"
    return (
        project_root
        / config.output_directory
        / str(config.year)
        / f"{config.month:02d}"
        / week_folder
    )


def format_date(value: date) -> str:
    """Format one title date in the established day/month/year style."""
    return f"{value.day}/{value.month}/{value.year}"
